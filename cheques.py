import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


CARPETA_PROGRAMA = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parent
)
ARCHIVO_EXCEL = CARPETA_PROGRAMA / "cheques.xlsx"
ARCHIVO_CSV_ANTERIOR = CARPETA_PROGRAMA / "cheques.csv"
COLUMNAS = ("numero", "fecha", "emisor", "banco", "monto", "concepto")
TITULOS = ("Número de cheque", "Fecha de recepción", "Emisor", "Banco", "Monto", "Concepto")
HOJA_HISTORIAL = "Historial"
HOJA_RESUMEN = "Resumen"


def crear_libro():
    libro = Workbook()
    historial = libro.active
    historial.title = HOJA_HISTORIAL
    historial.append(TITULOS)
    historial.freeze_panes = "A2"
    historial.auto_filter.ref = "A1:F1"

    for celda in historial[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")

    for columna, ancho in {"A": 20, "B": 22, "C": 28, "D": 22, "E": 16, "F": 36}.items():
        historial.column_dimensions[columna].width = ancho

    resumen = libro.create_sheet(HOJA_RESUMEN)
    resumen["A1"] = "Resumen de cheques recibidos"
    resumen["A1"].font = Font(bold=True, size=14, color="1F4E78")
    resumen["A3"] = "Cantidad de cheques"
    resumen["B3"] = f"=COUNTA('{HOJA_HISTORIAL}'!A2:A1048576)"
    resumen["A4"] = "Total recibido"
    resumen["B4"] = f"=SUM('{HOJA_HISTORIAL}'!E2:E1048576)"
    resumen["B4"].number_format = "#,##0.00"
    resumen.column_dimensions["A"].width = 28
    resumen.column_dimensions["B"].width = 20

    libro.calculation.fullCalcOnLoad = True
    libro.calculation.forceFullCalc = True
    return libro


def formatear_fila(historial, fila):
    historial.cell(row=fila, column=2).number_format = "yyyy-mm-dd"
    historial.cell(row=fila, column=5).number_format = "#,##0.00"


def importar_csv_anterior(historial):
    if not ARCHIVO_CSV_ANTERIOR.exists():
        return

    with ARCHIVO_CSV_ANTERIOR.open("r", newline="", encoding="utf-8") as archivo:
        for cheque in csv.DictReader(archivo):
            historial.append(
                (
                    cheque["numero"],
                    date.fromisoformat(cheque["fecha"]),
                    cheque["emisor"],
                    cheque["banco"],
                    float(Decimal(cheque["monto"])),
                    cheque["concepto"],
                )
            )
            formatear_fila(historial, historial.max_row)

    historial.auto_filter.ref = f"A1:F{historial.max_row}"


def pedir_texto(mensaje):
    while True:
        texto = input(mensaje).strip()
        if texto:
            return texto
        print("Este campo no puede quedar vacío.")


def pedir_fecha():
    while True:
        texto = input("Fecha de recepción (AAAA-MM-DD): ").strip()
        try:
            return date.fromisoformat(texto).isoformat()
        except ValueError:
            print("Fecha inválida. Usa el formato AAAA-MM-DD.")


def pedir_monto():
    while True:
        texto = input("Monto del cheque: ").strip().replace(",", ".")
        try:
            monto = Decimal(texto)
            if monto.is_finite() and monto > 0:
                return monto
        except InvalidOperation:
            pass
        print("Monto inválido. Ingresa un número mayor que cero.")


def registrar_cheque():
    cheque = {
        "numero": pedir_texto("Número del cheque: "),
        "fecha": date.fromisoformat(pedir_fecha()),
        "emisor": pedir_texto("Nombre de quien emite el cheque: "),
        "banco": pedir_texto("Banco: "),
        "monto": float(pedir_monto()),
        "concepto": pedir_texto("Concepto o motivo: "),
    }

    if ARCHIVO_EXCEL.exists():
        libro = load_workbook(ARCHIVO_EXCEL)
    else:
        libro = crear_libro()
        importar_csv_anterior(libro[HOJA_HISTORIAL])

    historial = libro[HOJA_HISTORIAL]
    historial.append(tuple(cheque[columna] for columna in COLUMNAS))
    fila = historial.max_row
    formatear_fila(historial, fila)
    historial.auto_filter.ref = f"A1:F{fila}"
    libro.save(ARCHIVO_EXCEL)

    print("Cheque registrado correctamente.")


def leer_cheques():
    if not ARCHIVO_EXCEL.exists():
        libro = crear_libro()
        importar_csv_anterior(libro[HOJA_HISTORIAL])
        libro.save(ARCHIVO_EXCEL)

    libro = load_workbook(ARCHIVO_EXCEL, data_only=True, read_only=True)
    try:
        hoja = libro[HOJA_HISTORIAL]
        return [
            dict(zip(COLUMNAS, valores))
            for valores in hoja.iter_rows(min_row=2, values_only=True)
            if any(valor is not None for valor in valores)
        ]
    finally:
        libro.close()


def formatear_fecha(fecha):
    if isinstance(fecha, datetime):
        return fecha.date().isoformat()
    return fecha.isoformat() if isinstance(fecha, date) else str(fecha)


def mostrar_cheques():
    cheques = leer_cheques()
    if not cheques:
        print("Todavía no hay cheques registrados.")
        return

    print("\nCheques recibidos:")
    for cheque in cheques:
        print(
            f"N.º {cheque['numero']} | {formatear_fecha(cheque['fecha'])} | "
            f"{cheque['emisor']} | {cheque['banco']} | "
            f"Monto: {Decimal(str(cheque['monto'])):.2f} | {cheque['concepto']}"
        )


def mostrar_total():
    cheques = leer_cheques()
    total = sum(
        (Decimal(str(cheque["monto"])) for cheque in cheques),
        Decimal("0"),
    )
    print(f"Total de cheques recibidos: {total:.2f}")


def main():
    while True:
        print("\n=== Registro de cheques recibidos ===")
        print("1. Registrar un cheque")
        print("2. Ver cheques registrados")
        print("3. Ver total recibido")
        print("4. Salir")
        opcion = input("Elige una opción: ").strip()

        if opcion == "1":
            registrar_cheque()
        elif opcion == "2":
            mostrar_cheques()
        elif opcion == "3":
            mostrar_total()
        elif opcion == "4":
            print("Hasta luego.")
            break
        else:
            print("Opción no válida. Elige del 1 al 4.")


if __name__ == "__main__":
    main()
